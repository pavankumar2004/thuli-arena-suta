"""Build SUTA_Digital_Lookbook_Presentation.pptx from presentation_assets/ screenshots.

    uv run --no-project --with python-pptx --with pillow python build_deck.py

Screenshots come from a headless run of the local app (npm run dev in lookbook/).
Every number on the slides is taken from the code, data/export/REPORT.md or a measured run.
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).parent
ASSETS = ROOT / "presentation_assets"
CROPS = ASSETS / "_crops"
OUT = ROOT / "SUTA_Digital_Lookbook_Presentation.pptx"

ECRU, KORA, INK = "FAF7F2", "F1EBE1", "1C2A3A"
RUST, HALDI, CHAR, SLATE = "A4422D", "D99B26", "1A1A1A", "5A626A"
LINE, ECRU_DIM = "D9D0C2", "B9C0C8"
SERIF, SANS = "Georgia", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
used_assets: list[str] = []


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# ---------- primitives ----------

def new_slide(bg: str = ECRU, notes: str = ""):
    s = prs.slides.add_slide(BLANK)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(bg)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, runs, *, size=14, color=CHAR, font=SANS, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line=None, after=0):
    """runs: a string, or a list of paragraphs; each paragraph a string or list of (text, overrides)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        p.space_after = Pt(after)
        for chunk in (para if isinstance(para, list) else [(para, {})]):
            t, o = chunk if isinstance(chunk, tuple) else (chunk, {})
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
            sp = o.get("spacing", spacing)
            if sp is not None:
                r._r.get_or_add_rPr().set("spc", str(int(sp * 100)))
    return tb


def eyebrow(s, x, y, label, color=RUST, w=8):
    return text(s, x, y, w, 0.3, label.upper(), size=10, color=color, bold=True, spacing=3)


def rect(s, x, y, w, h, fill=None, line=None, weight=0.75, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        r.fill.solid()
        r.fill.fore_color.rgb = rgb(fill)
    else:
        r.fill.background()
    if line:
        r.line.color.rgb = rgb(line)
        r.line.width = Pt(weight)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def rule(s, x, y, w, color=LINE, weight=0.75):
    """A hairline drawn as a filled rectangle (connectors pick up the theme's shadow)."""
    return rect(s, x, y, w, max(weight, 0.5) / 72, fill=color)


def vrule(s, x, y, h, color=LINE, weight=0.75):
    return rect(s, x, y, max(weight, 0.5) / 72, h, fill=color)


def arrow(s, x1, y1, x2, y2, color=SLATE):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    sp = c._element.spPr
    sp.append(sp.makeelement(qn("a:effectLst"), {}))
    return c


def pill(s, x, y, label, *, fill=None, line=RUST, color=RUST, size=10, pad=0.18, h=0.32):
    w = pad * 2 + len(label) * size * 0.0075
    r = rect(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    r.adjustments[0] = 0.5
    tf = r.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size, run.font.name, run.font.color.rgb = Pt(size), SANS, rgb(color)
    return x + w


def crop(name, box=None, aspect=None, anchor=(0.5, 0.5), tag=""):
    """Crop a screenshot (optional pixel box), then centre-crop to `aspect` (w/h)."""
    src = Path(name) if Path(name).is_absolute() else ASSETS / name
    if name not in used_assets:
        used_assets.append(name)
    im = Image.open(src).convert("RGB")
    if box:
        im = im.crop(box)
    if aspect:
        w, h = im.size
        if w / h > aspect:
            nw = int(h * aspect)
            left = int((w - nw) * anchor[0])
            im = im.crop((left, 0, left + nw, h))
        else:
            nh = int(w / aspect)
            top = int((h - nh) * anchor[1])
            im = im.crop((0, top, w, top + nh))
    CROPS.mkdir(exist_ok=True)
    out = CROPS / f"{Path(name).stem}{tag}_{im.size[0]}x{im.size[1]}.png"
    im.save(out)
    return out


def picture(s, path, x, y, w, h, frame=LINE, weight=0.75):
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if frame:
        pic.line.color.rgb = rgb(frame)
        pic.line.width = Pt(weight)
    return pic


def shot(s, name, x, y, w, h, box=None, anchor=(0.5, 0.5), frame=LINE, tag=""):
    return picture(s, crop(name, box, w / h, anchor, tag), x, y, w, h, frame)


def footer(s, n, dark=False):
    c = ECRU_DIM if dark else SLATE
    text(s, 0.6, 7.05, 6, 0.25, "SUTA  ·  THULI ARENA", size=8, color=c, bold=True, spacing=3)
    text(s, 11.73, 7.05, 1.0, 0.25, f"{n:02d}", size=8, color=c, bold=True, spacing=2, align=PP_ALIGN.RIGHT)


def title(s, x, y, w, t, size=34, color=CHAR, h=1.2):
    return text(s, x, y, w, h, t, size=size, color=color, font=SERIF, line=1.0)


def bullets(s, x, y, w, h, items, size=13, color=CHAR, gap=7, mark=RUST):
    paras = []
    for it in items:
        head, _, rest = it.partition("|")
        para = [("—  ", {"color": mark, "bold": True})]
        if rest:
            para += [(head, {"bold": True}), (" " + rest, {})]
        else:
            para += [(head, {})]
        paras.append(para)
    return text(s, x, y, w, h, paras, size=size, color=color, line=1.12, after=gap)


def stat(s, x, y, w, number, label, ncolor=CHAR, lcolor=SLATE, nsize=34):
    text(s, x, y, w, 0.65, number, size=nsize, color=ncolor, font=SERIF)
    text(s, x, y + nsize / 72 * 1.2 + 0.04, w, 0.5, label, size=10.5, color=lcolor, line=1.05)


# ---------- slides ----------
n = 0


def next_n():
    global n
    n += 1
    return n


# 1 · Cover
s = new_slide(INK, "Opening. We built four connected things on one source of truth: Suta's live catalogue.")
shot(s, "A1_cover.png", 6.9, 0, 6.433, 7.5, anchor=(0.5, 0.5), frame=None)
rect(s, 6.9, 0, 0.04, 7.5, fill=HALDI)
eyebrow(s, 0.7, 0.75, "Thuli Arena  ·  Hackathon submission", color=HALDI)
title(s, 0.7, 1.55, 5.9, "Suta, woven\ninto software", size=54, color=ECRU, h=2.4)
text(s, 0.7, 3.95, 5.6, 1.4,
     "A verified catalogue, an editorial lookbook, a lookbook made for one person, and a stylist "
     "with a fitting room. Four builds, one rule: every piece shown is real, at its real price.",
     size=15, color="D7DCE2", line=1.2)
rule(s, 0.7, 5.75, 5.6, color="3A4A5C")
for i, (k, v) in enumerate([("Task 1", "Catalogue"), ("Task 2", "Lookbook"), ("Task 3", "Made for one"), ("Task 4", "Stylist + try-on")]):
    text(s, 0.7 + i * 1.45, 5.95, 1.4, 0.6, [[(k.upper(), {"size": 8.5, "color": HALDI, "bold": True, "spacing": 2})],
                                            [(v, {"size": 12, "color": ECRU, "font": SERIF})]], line=1.2)
next_n()

# 2 · Brand & vision
s = new_slide(ECRU, "Suta's voice is nostalgic and personal. Our build had to keep that voice and never invent a product.")
eyebrow(s, 0.7, 0.7, "01  ·  The brand")
title(s, 0.7, 1.1, 6.0, "Handloom with\na memory", size=44, h=1.9)
text(s, 0.7, 3.15, 5.5, 2.4,
     [[("Suta makes handloom sarees, blouses and everyday Indian wear that feel like they came out of a "
        "grandmother's trunk: soft mul cotton, real weaves, and a story behind every product name.", {})],
      [("", {})],
      [("Our job was to bring that soul online without losing a single fact: the same nostalgia in the "
        "pages, and the catalogue's truth underneath every word.", {"color": SLATE})]],
     size=15, line=1.25)
shot(s, "A2_chapter_switcher.png", 7.0, 0.7, 5.65, 2.95, box=(0, 65, 1920, 905), anchor=(0.35, 0.5))
text(s, 7.0, 3.75, 5.65, 0.3, "Chapter I · Agomoni, from the live lookbook", size=9.5, color=SLATE, italic=True)
principles = [("Grounded", "Nothing invented. Every card, price and size is a database row."),
              ("Editorial", "Restraint over clutter: rhythm, whitespace, the brand's own photography."),
              ("Fast", "Cards in under 2 s; the stylist starts writing in under 3 s.")]
for i, (h, b) in enumerate(principles):
    y = 4.35 + i * 0.82
    text(s, 7.0, y, 0.5, 0.5, f"0{i + 1}", size=20, color=HALDI, font=SERIF)
    text(s, 7.6, y + 0.02, 5.0, 0.8, [[(h, {"bold": True, "size": 13})], [(b, {"color": SLATE, "size": 12})]], line=1.1)
footer(s, next_n())

# 3 · Architecture
s = new_slide(ECRU, "One pipeline feeds everything. The catalogue is scraped and verified once; three experiences read it.")
eyebrow(s, 0.7, 0.6, "02  ·  Architecture")
title(s, 0.7, 0.95, 12, "One source of truth, three experiences", size=32, h=0.8)


def box(x, y, w, h, head, body, fill=KORA, hc=CHAR, bc=SLATE, line=LINE, eb=None, ebc=RUST):
    rect(s, x, y, w, h, fill=fill, line=line)
    yy = y + 0.14
    if eb:
        text(s, x + 0.18, yy, w - 0.3, 0.25, eb.upper(), size=8, color=ebc, bold=True, spacing=2)
        yy += 0.27
    text(s, x + 0.18, yy, w - 0.3, 0.4, head, size=14, color=hc, font=SERIF)
    text(s, x + 0.18, yy + 0.38, w - 0.3, h - (yy - y) - 0.45, body, size=10, color=bc, line=1.1)


box(0.7, 2.15, 2.2, 1.55, "suta.in", "Shopify storefront's public JSON: products, variants, collections. robots.txt-aware, 1 req/s.", eb="Source")
arrow(s, 2.9, 2.92, 3.3, 2.92)
box(3.3, 2.15, 2.75, 1.55, "Catalogue pipeline", "Python: fetch → scrape → build → normalise → export. Checks against the live site.", eb="Task 1")
arrow(s, 6.05, 2.92, 6.45, 2.92)
box(6.45, 1.95, 2.6, 1.95, "Neon Postgres", "7,249 products · 19,246 variants. pg_trgm + GIN indexes on tags, edits, sizes, colours.",
    fill=INK, hc=ECRU, bc="C9D0D8", line=INK, eb="One source of truth", ebc=HALDI)
consumers = [("Utsav lookbook", "Pre-rendered pages; build fails if a look names a product not in the catalogue.", "Task 2", 1.55),
             ("Made for one", "Instagram posts → taste → SQL shortlist → three styled looks.", "Task 3", 3.15),
             ("Stylist + try-on", "Node API routes: planner, code-built SQL, streamed writer, image model.", "Task 4", 4.75)]
for head, body, eb, y in consumers:
    box(9.65, y, 3.0, 1.38, head, body, eb=eb, fill=ECRU)
    arrow(s, 9.05, 2.92, 9.65, y + 0.69)
rect(s, 0.7, 4.3, 8.35, 2.35, fill=None, line=LINE)
text(s, 0.95, 4.45, 4, 0.3, "MODELS, VIA ONE OPENROUTER CLIENT", size=8, color=RUST, bold=True, spacing=2)
models = [("GPT-4.1 mini", "Stylist planner + writer, photo search, fit guard"),
          ("Gemini 2.5 Flash", "Tags Instagram + catalogue photos"),
          ("Claude Sonnet 5.5", "Reads taste, styles the three looks"),
          ("Gemini 3.1 Flash Image", "On-demand virtual try-on"),
          ("Gemini 3 Pro Image", "Editorial lookbook imagery")]
for i, (m_, use) in enumerate(models):
    col, row = i % 2, i // 2
    text(s, 0.95 + col * 4.05, 4.85 + row * 0.58, 3.9, 0.55,
         [[(m_, {"bold": True, "size": 11.5})], [(use, {"color": SLATE, "size": 10})]], line=1.05)
text(s, 0.95 + 4.05, 4.85 + 2 * 0.58, 3.9, 0.55,
     [[("Shared client", {"bold": True, "size": 11.5})], [("Retries, fallbacks, cost logged per run", {"color": SLATE, "size": 10})]], line=1.05)
footer(s, next_n())

# 4 · Stack & why
s = new_slide(ECRU, "Each choice is justified by latency, grounding, scale or aesthetics.")
eyebrow(s, 0.7, 0.6, "03  ·  Stack, and why")
title(s, 0.7, 0.95, 12, "Chosen for speed, truth and taste", size=32, h=0.8)
rows = [("Data", "Python · Shopify public JSON", "Exact prices, variants and stock without parsing HTML; a full cold run takes about 4 minutes and is repeatable."),
        ("Database", "Neon serverless Postgres", "One catalogue for every task. Search is structured SQL over real attributes (no embeddings), so every match is explainable."),
        ("Speed", "WebSocket pool · hedged reads", "A warm pool, a second read fired at 1.2 s and a 5-minute cache keep the stylist's queries fast."),
        ("App", "Next.js 16 App Router · React 19", "Pre-rendered lookbook pages and streaming AI routes in one codebase; secrets stay server-side."),
        ("Models", "OpenRouter", "The right model per job, a fallback per call, and cost logged per run."),
        ("Interface", "Tailwind v4 · shadcn/ui (Radix)", "Brand tokens and accessible dialogs, sheets and focus handling."),
        ("Motion", "Framer Motion (lazy) · Lenis", "Editorial pacing without a heavy bundle; smooth scroll only for mouse and trackpad.")]
y0 = 1.95
text(s, 0.7, y0, 1.6, 0.3, "LAYER", size=8.5, color=SLATE, bold=True, spacing=2)
text(s, 2.4, y0, 3.6, 0.3, "CHOICE", size=8.5, color=SLATE, bold=True, spacing=2)
text(s, 6.2, y0, 6.4, 0.3, "WHY IT MATTERS", size=8.5, color=SLATE, bold=True, spacing=2)
rule(s, 0.7, y0 + 0.33, 11.93, color=CHAR, weight=1)
for i, (layer, choice, why) in enumerate(rows):
    y = y0 + 0.45 + i * 0.64
    text(s, 0.7, y, 1.6, 0.55, layer, size=12, color=RUST, bold=True)
    text(s, 2.4, y - 0.02, 3.7, 0.55, choice, size=13.5, font=SERIF)
    text(s, 6.2, y, 6.4, 0.6, why, size=11, color=SLATE, line=1.08)
    rule(s, 0.7, y + 0.56, 11.93)
footer(s, next_n())

# 5 · Task 1 metrics
s = new_slide(ECRU, "Task 1 numbers come straight from data/export/REPORT.md.")
eyebrow(s, 0.7, 0.6, "Task 1  ·  The catalogue")
title(s, 0.7, 0.95, 12, "Every product, exactly as suta.in sells it", size=32, h=0.8)
stats = [("7,249", "products exported\n(9 add-ons and test items excluded)"),
         ("19,246", "variants: each size with its\nown price and stock flag"),
         ("49,288", "full-size images,\nfrom Shopify's CDN"),
         ("37 / 37", "categories match the counts\nsuta.in itself shows")]
for i, (num, lab) in enumerate(stats):
    x = 0.7 + i * 3.05
    rule(s, x, 2.05, 2.75, color=CHAR, weight=1.25)
    stat(s, x, 2.2, 2.8, num, lab, nsize=40)
text(s, 0.7, 3.85, 6, 0.3, "FIELD COMPLETENESS", size=8.5, color=SLATE, bold=True, spacing=2)
x = 0.7
for f in ["title 100%", "price 100%", "images 100%", "sizes 100%", "category 100%", "url 100%", "colours 99.9%"]:
    x = pill(s, x, 4.2, f, size=10.5) + 0.12
text(s, 0.7, 4.95, 6, 0.3, "THE PIPELINE", size=8.5, color=SLATE, bold=True, spacing=2)
steps = [("fetch.py", "robots-aware, 1 req/s, retries, disk cache"), ("scrape.py", "products + every menu collection"),
         ("build.py", "drop non-products, attach categories"), ("normalise.py", "one clean record per product"),
         ("export · db", "JSONL + REPORT.md + Neon")]
for i, (h, b) in enumerate(steps):
    x = 0.7 + i * 2.42
    rect(s, x, 5.3, 2.18, 1.2, fill=KORA, line=LINE)
    text(s, x + 0.15, 5.42, 1.95, 0.35, h, size=12.5, font=SERIF)
    text(s, x + 0.15, 5.78, 1.95, 0.7, b, size=9.5, color=SLATE, line=1.08)
    if i < 4:
        arrow(s, x + 2.18, 5.9, x + 2.42, 5.9)
footer(s, next_n())

# 6 · Task 1 design decisions
s = new_slide(ECRU, "Design decisions that make the data trustworthy, plus the fields later tasks needed.")
eyebrow(s, 0.7, 0.6, "Task 1  ·  Data integrity")
title(s, 0.7, 0.95, 7.5, "Built to survive a judge\nre-checking the live site", size=30, h=1.2)
decisions = [("Categories from the menu", "Not Shopify's free-text product_type: judges compare against the brand's own category counts."),
             ("Nothing missed mid-scrape", "Every category feed is merged and de-duplicated by Shopify id, keeping the newest copy."),
             ("Sizes are variants", "Nested under their product, with per-size stock in sizes_in_stock."),
             ("Colours, most reliable first", "Colour_ tags (97%), then title/URL in English or Hindi, then the photo itself (198 products)."),
             ("Fields for later tasks", "fabric, pattern, style, technique facets; the brand's edits; pairs_with from its own copy."),
             ("Verified, not assumed", "catalogue.check: unique ids, Neon matches the export, and 20 random products re-fetched live.")]
for i, (h, b) in enumerate(decisions):
    col, row = i % 2, i // 2
    x, y = 0.7 + col * 3.95, 2.45 + row * 1.42
    rule(s, x, y, 3.65, color=RUST, weight=1)
    text(s, x, y + 0.12, 3.65, 1.25, [[(h, {"font": SERIF, "size": 15})], [(b, {"size": 10.5, "color": SLATE})]], line=1.12, after=3)
rect(s, 8.75, 0.95, 3.9, 5.75, fill=INK)
text(s, 9.0, 1.15, 3.5, 0.3, "PRODUCTS  ·  KEY COLUMNS", size=8, color=HALDI, bold=True, spacing=2)
schema = [("id", "bigint, Shopify id"), ("handle · url · title", "text"), ("price · compare_at_price", "numeric"),
          ("available", "boolean"), ("sizes · sizes_in_stock", "text[]"), ("colours · colour_source", "text[] · text"),
          ("department · category", "text"), ("categories · edits", "text[]"), ("attributes", "jsonb facets"),
          ("tags", "text[]"), ("pairs_with", "text[]"), ("images", "text[]"), ("description", "text")]
paras = []
for c, t in schema:
    paras.append([(c, {"font": "Consolas", "size": 10.5, "color": ECRU}), ("   " + t, {"font": "Consolas", "size": 9, "color": "8E9AA8"})])
text(s, 9.0, 1.55, 3.55, 5.0, paras, line=1.0, after=6)
footer(s, next_n())

# 7 · Task 2 hero
s = new_slide(ECRU, "The lookbook: six chapters on six of Suta's own collections. All prices from the catalogue export.")
eyebrow(s, 0.7, 0.6, "Task 2  ·  The lookbook")
title(s, 0.7, 0.95, 5.3, "Utsav: a festive\nalmanac in six\nchapters", size=38, h=2.2)
text(s, 0.7, 3.25, 5.0, 1.3,
     "One festive season, from Durga Puja to the quiet after Diwali, told the way a fashion house prints "
     "its lookbook: one idea per spread, the brand's own photography, and quiet type.", size=13, color=SLATE, line=1.2)
x, y = 0.7, 4.55
for i, c in enumerate(["I Agomoni", "II Zar", "III Chandrama", "IV Jalebi", "V Varq", "VI Kavita"]):
    if i == 3:
        x, y = 0.7, 4.95
    x = pill(s, x, y, c, size=10) + 0.1
for i, (num, lab) in enumerate([("24", "looks"), ("42", "real pieces"), ("16", "new editorial\nimages")]):
    stat(s, 0.7 + i * 1.7, 5.55, 1.6, num, lab, nsize=30)
shot(s, "A1_cover.png", 6.35, 0.6, 6.3, 6.1)
text(s, 6.35, 6.75, 6.3, 0.25, "The cover, live at localhost:3000", size=9, color=SLATE, italic=True)
footer(s, next_n())

# 8 · Task 2 interactions
s = new_slide(ECRU, "Shoppable hotspots, the look drawer, and build-time grounding.")
eyebrow(s, 0.7, 0.6, "Task 2  ·  From image to product")
title(s, 0.7, 0.95, 12, "Tap the photograph, reach the real piece", size=30, h=0.8)
shot(s, "B1_hotspot_figure.png", 0.7, 1.9, 3.55, 4.2, anchor=(0.5, 0.3))
shot(s, "B2_look_drawer.png", 4.5, 1.9, 5.1, 4.2, box=(880, 0, 1920, 1080), anchor=(0.5, 0.0))
text(s, 0.7, 6.18, 3.55, 0.6, "Hotspot card: piece, fabric, sale price, link. It always stays inside the frame.", size=9.5, color=SLATE, italic=True, line=1.1)
text(s, 4.5, 6.18, 5.1, 0.6, "Look drawer: weave zoom, the label's own story, sizes in stock, the total for the look.", size=9.5, color=SLATE, italic=True, line=1.1)
bullets(s, 9.9, 1.95, 2.8, 4.8, [
    "Grounded at build time.|The build fails if a look names a product that isn't in the catalogue.",
    "Nothing hand-typed.|Titles, prices, stock, sizes and fabric all come from the export.",
    "Three clicks|from a look to suta.in.",
    "Fast.|Lighthouse desktop 99–100, mobile 83–89; accessibility, SEO and best practices 100.",
], size=11.5, gap=9)
footer(s, next_n())

# 9 · Task 3
s = new_slide(ECRU, "Made for one: public Instagram posts become three real looks, each tied to a post.")
eyebrow(s, 0.7, 0.6, "Task 3  ·  Made for one")
title(s, 0.7, 0.95, 5.6, "Three looks, styled\nfrom your own posts", size=32, h=1.3)
flow = [("Photos", "Public posts saved to Neon once; repeat visits skip Instagram."),
        ("Tag", "A fast vision model labels each outfit in the catalogue's own vocabulary."),
        ("Taste", "A strong vision model reads style, palette and three moments from life."),
        ("Shortlist", "No AI here: SQL picks in-stock pieces that match each moment."),
        ("Style", "One call picks one piece per moment and writes why. Every id is validated.")]
for i, (h, b) in enumerate(flow):
    y = 2.35 + i * 0.68
    rect(s, 0.7, y + 0.04, 0.34, 0.34, fill=RUST if i != 3 else INK, shape=MSO_SHAPE.OVAL)
    text(s, 0.7, y + 0.06, 0.34, 0.3, str(i + 1), size=10, color=ECRU, bold=True, align=PP_ALIGN.CENTER)
    text(s, 1.2, y, 4.9, 0.7, [[(h, {"bold": True, "size": 12.5})], [(b, {"size": 10.5, "color": SLATE})]], line=1.05)
for i, (num, lab) in enumerate([("4–18 s", "saved profile"), ("~30–40 s", "first-time profile"), ("$0.02–0.06", "per lookbook")]):
    stat(s, 0.7 + i * 1.85, 5.95, 1.8, num, lab, nsize=19)
shot(s, "C5_made_for_one_look.png", 6.45, 0.6, 6.2, 6.1, box=(436, 44, 1484, 1080), anchor=(0.5, 0.0))
text(s, 6.45, 6.75, 6.2, 0.25, "A finished lookbook: their post on the left, the reasoning and real pieces on the right", size=9, color=SLATE, italic=True)
footer(s, next_n())

# 10 · Task 4 stylist
s = new_slide(ECRU, "The stylist: rules and a planner decide what to fetch; code builds the SQL; a writer streams a checked reply.")
shot(s, "D1_stylist_top5.png", 0.7, 0.6, 5.3, 6.4, box=(1024, 0, 1920, 1080), anchor=(0.5, 0.0))
eyebrow(s, 6.5, 0.6, "Task 4  ·  The stylist")
title(s, 6.5, 0.95, 6.2, "A real conversation,\ngrounded in Neon", size=32, h=1.3)
pipe = [("Guard", "size caps, injection stripped, rate limits"),
        ("Rules + planner", "instant parser and GPT-4.1 mini choose filters from the catalogue's own values"),
        ("Code-built SQL", "one parameterised query ranks pieces by criteria met"),
        ("Cards first", "the top five real products stream to the panel"),
        ("Writer", "streams the reply; every sentence checked before it is sent")]
for i, (h, b) in enumerate(pipe):
    y = 2.5 + i * 0.66
    rect(s, 6.5, y, 0.08, 0.5, fill=HALDI if i in (2, 3) else RUST)
    text(s, 6.75, y - 0.02, 5.9, 0.62, [[(h, {"bold": True, "size": 12.5})], [(b, {"size": 10.5, "color": SLATE})]], line=1.05)
rule(s, 6.5, 5.95, 6.15)
for i, (num, lab) in enumerate([("1.2–1.9 s", "cards on screen"), ("2.2–2.9 s", "first words of the reply"), ("≤ 350", "max tokens, any call")]):
    stat(s, 6.5 + i * 2.1, 6.05, 2.0, num, lab, nsize=22)
footer(s, next_n())

# 11 · Honesty + curveball
s = new_slide(ECRU, "Honesty circuit-breakers, session memory, and the 15:45 curveball occasion filter.")
eyebrow(s, 0.7, 0.6, "Task 4  ·  Honesty circuit-breakers")
title(s, 0.7, 0.95, 12, "It would rather say no than invent", size=32, h=0.8)
cols = [("Never invents", ["The model never writes SQL; filters come only from catalogue values.",
                           "Any ₹ amount in a reply must be a price we fetched.",
                           "A failed check swaps in a plain line built from the same rows."]),
        ("Honest when short", ["Nothing matches? One query finds what does exist nearby, with real counts.",
                               "“Festive sarees under ₹2,000 · 12” becomes a tappable suggestion.",
                               "Near misses carry a “Close match” label."]),
        ("Remembers", ["“Something else” leaves out everything already shown this session.",
                       "Tested: 4 turns, 20 distinct greens, 0 repeats.",
                       "Survives a reload; “New chat” clears it."])]
for i, (h, items) in enumerate(cols):
    x = 0.7 + i * 4.05
    rule(s, x, 1.95, 3.75, color=CHAR, weight=1.25)
    text(s, x, 2.08, 3.75, 0.4, h, size=17, font=SERIF)
    bullets(s, x, 2.6, 3.75, 2.2, items, size=11, gap=6)
rect(s, 0.7, 4.6, 11.95, 2.15, fill=INK)
text(s, 1.0, 4.85, 4.0, 0.3, "THE 15:45 CURVEBALL", size=8.5, color=HALDI, bold=True, spacing=2)
text(s, 1.0, 5.25, 5.0, 1.3, "Occasion edit: one tap sets a hard filter on the catalogue's style and occasion tags, "
     "and it stacks with whatever the shopper types.", size=13, color=ECRU, font=SERIF, line=1.2)
picture(s, crop("D1_stylist_top5.png", (1040, 82, 1556, 126), tag="_chips"), 6.4, 5.2, 6.0, 0.51, frame=None)
text(s, 6.4, 5.88, 6.0, 0.6, "Asked for “a pure Kanjeevaram bridal saree under ₹10”, it says plainly that no such piece exists, "
     "then offers what does.", size=10, color="C9D0D8", italic=True, line=1.15)
footer(s, next_n())

# 12 · Try-on
s = new_slide(ECRU, "Try-on happens only when asked, for one piece. Measured 20 seconds in our test.")
eyebrow(s, 0.7, 0.6, "Task 4  ·  The fitting room")
title(s, 0.7, 0.95, 5.4, "One click, one\ntry-on, on demand", size=34, h=1.3)
bullets(s, 0.7, 2.45, 5.2, 3.4, [
    "Only the piece you choose.|No batch generation, so no timeouts and no wasted spend.",
    "Faithful to the weave.|Your portrait plus the product's own catalogue photos go to Gemini 3.1 Flash Image.",
    "Before and after.|Drag the line to compare; the result is saved in the chat to download later.",
    "Bounded.|4 per visitor per 10 minutes, 3 at once; identical requests served from cache.",
    "Private.|Portraits are re-encoded with EXIF stripped and never stored on the server.",
], size=11.5, gap=8)
stat(s, 0.7, 6.0, 2.5, "~20 s", "measured, portrait to finished image", nsize=24)
shot(s, "E3_tryon_split.png", 6.3, 0.6, 6.35, 4.53, box=(576, 266, 1344, 814), tag="_modal")
shot(s, "E1_tryon_working.png", 6.3, 5.3, 3.05, 1.55, box=(576, 266, 987, 814), anchor=(0.5, 0.55), tag="_modal")
text(s, 9.55, 5.35, 3.1, 1.5, "While it works, the fitting room shows its progress: reading the photo, studying the weave, "
     "pleating the drape, setting the light.", size=10, color=SLATE, italic=True, line=1.15)
footer(s, next_n())

# 13 · Hardening
s = new_slide(ECRU, "Task 5 readiness: abuse is handled with plain 4xx answers, never 500s, and costs are capped.")
eyebrow(s, 0.7, 0.6, "Ready for the stress test")
title(s, 0.7, 0.95, 7.2, "Hardened against misuse", size=32, h=0.8)
guards = [("Prompt injection", "Known patterns stripped and flagged; product and caption text treated as data."),
          ("Token burning", "Over 500 characters rejected; messages capped at 350; max_tokens ≤ 350."),
          ("Malformed input", "Empty, emoji-only or wrongly typed bodies get a 400 with a kind answer."),
          ("Images", "Data URLs only (no URL fetching, no SSRF); 6 MB cap; re-encoded."),
          ("Rate limits", "Chat 30 a minute; try-on 4 per 10 minutes and 3 at once."),
          ("Leaks", "Replies naming prompts, keys or the database are discarded; secrets stay server-side."),
          ("Fit guardrail", "No womenswear try-on on a man's photo, or menswear on a woman's; unisex passes."),
          ("Traceable", "In development, “Behind this reply” shows each Neon query, its timing and every model prompt.")]
for i, (h, b) in enumerate(guards):
    col, row = i % 2, i // 2
    x, y = 0.7 + col * 3.6, 1.95 + row * 1.2
    rule(s, x, y, 3.35, color=RUST if h != "Fit guardrail" else HALDI, weight=1)
    text(s, x, y + 0.1, 3.35, 1.1, [[(h, {"font": SERIF, "size": 14})], [(b, {"size": 10.5, "color": SLATE})]], line=1.1, after=3)
shot(s, "E5_guardrail_refusal.png", 8.2, 1.95, 4.45, 3.18, box=(576, 266, 1344, 814), tag="_modal")
text(s, 8.2, 5.25, 4.45, 1.0, "The fit guardrail at work: the catalogue says the saree is womenswear, a vision check reads the photo, "
     "and nothing is generated.", size=10, color=SLATE, italic=True, line=1.15)
footer(s, next_n())

# 14 · Edge
s = new_slide(ECRU, "Where the design differs from the usual approach.")
eyebrow(s, 0.7, 0.6, "What sets it apart")
title(s, 0.7, 0.95, 12, "Design choices that hold up under pressure", size=32, h=0.8)
text(s, 0.7, 1.95, 5.6, 0.3, "THE USUAL APPROACH", size=8.5, color=SLATE, bold=True, spacing=2)
text(s, 6.9, 1.95, 5.8, 0.3, "WHAT WE BUILT", size=8.5, color=RUST, bold=True, spacing=2)
rule(s, 0.7, 2.3, 11.95, color=CHAR, weight=1.25)
pairs = [("An LLM answers from memory and names products", "The LLM only plans; code builds the SQL; every sentence is checked"),
         ("Try-ons generated for every card, slow and costly", "One try-on, on demand, cached, about 20 seconds"),
         ("Embedding search: plausible but unexplainable", "Structured SQL on real attributes; every match traceable"),
         ("“No results” dead ends", "Real alternatives with live counts, one tap away"),
         ("Mock or hand-typed product data", "A live catalogue, verified against suta.in"),
         ("Guardrails bolted on later", "Caps, rate limits, leak filter and fit guardrail from day one")]
for i, (a, b) in enumerate(pairs):
    y = 2.45 + i * 0.7
    text(s, 0.7, y, 5.6, 0.6, a, size=12.5, color=SLATE, line=1.05)
    text(s, 6.45, y - 0.02, 0.3, 0.5, "→", size=16, color=HALDI, bold=True)
    text(s, 6.9, y - 0.02, 5.75, 0.6, b, size=13.5, font=SERIF, line=1.05)
    rule(s, 0.7, y + 0.58, 11.95)
footer(s, next_n())

# 15 · Close
s = new_slide(INK, "Close. Run it with npm run dev inside lookbook/.")
shot(s, str(ROOT / "lookbook/public/images/looks/agomoni-01-shiuli-v2.jpg"), 0, 0, 5.0, 7.5, frame=None, anchor=(0.5, 0.3))
rect(s, 5.0, 0, 0.04, 7.5, fill=HALDI)
eyebrow(s, 5.8, 1.0, "In one line", color=HALDI)
title(s, 5.8, 1.45, 7, "Every piece real.\nEvery price true.", size=48, color=ECRU, h=2.2)
for i, (num, lab) in enumerate([("7,249", "products, verified"), ("24", "editorial looks"), ("3", "looks made for one"), ("5", "real cards per answer")]):
    stat(s, 5.8 + i * 1.75, 4.05, 1.7, num, lab, ncolor=ECRU, lcolor="B9C0C8", nsize=30)
rule(s, 5.8, 5.4, 6.85, color="3A4A5C")
text(s, 5.8, 5.6, 6.85, 0.9, [[("Run it:  ", {"color": HALDI, "bold": True}), ("cd lookbook  ·  npm run dev  ·  localhost:3000", {"font": "Consolas", "color": ECRU})],
                              [("Catalogue:  ", {"color": HALDI, "bold": True}), ("uv run python -m catalogue", {"font": "Consolas", "color": ECRU})]],
     size=12, line=1.4)
text(s, 5.8, 6.7, 6.85, 0.3, "Products and product photography © Suta; some editorial images generated for this build. An independent hackathon project, not affiliated with Suta.", size=8.5, color="8E9AA8")
next_n()

prs.save(OUT)
print(f"saved {OUT.name}: {len(prs.slides)} slides")
print("assets embedded:", ", ".join(used_assets))
